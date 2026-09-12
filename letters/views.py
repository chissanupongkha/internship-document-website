# -*- coding: utf-8 -*-
import io
import logging
import os
import re
import zipfile
from io import BytesIO
import subprocess
import tempfile
import zipfile
from pathlib import Path
from docx import Document
import pandas as pd
from django.conf import settings
from django.contrib import messages
from django.contrib.auth import login, logout, update_session_auth_hash
from django.contrib.auth.forms import AuthenticationForm, PasswordChangeForm
from django.core.files.base import ContentFile
from django.core.mail import EmailMessage
from django.http import FileResponse, Http404, HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.clickjacking import xframe_options_sameorigin

from .decorators import staff_required
from .docgen import (
    AVAILABLE_SOURCE_FIELDS,
    H_COMPANY,
    H_CONTACT_PERSON,
    H_CONTACT_POSITION,
    H_COURSE,
    H_FIRST_NAME,
    H_INTERN_POSITION,
    H_LAST_NAME,
    H_LETTER_TYPE,
    H_PERIOD,
    H_PREFIX,
    H_STUDENT_ID,
    H_YEAR,
    SUPPORTED_LETTER_TYPES,
    convert_docx_bytes_to_pdf,
    extract_fields_from_docx,
    generate_for_row,
    generate_pdf_for_row,
    normalize_letter_type,
    render_with_template,
    PDF_SAFE_FONT,
    _get_libreoffice_bin,
    apply_thai_font_fix,
    generate_for_row,
)
)
from .models import (
    AuditLog,
    GeneratedDocument,
    GenerationBatch,
    LetterTemplate,
    LetterType,
    LetterTypeMapping,
    StudentLetter,
    StudentRecord,
    TemplateFieldMapping,
    UploadedBatch,
)

IMPORTANT_FIELDS = ['student_id', 'display_name', 'company', 'contractor', 'internship_position', 'period', 'year']

raw_priority_headers = [
    H_STUDENT_ID, H_LETTER_TYPE, H_COURSE, H_COMPANY,
    H_CONTACT_PERSON, H_CONTACT_POSITION, H_INTERN_POSITION,
    H_PERIOD, H_PREFIX, H_FIRST_NAME, H_LAST_NAME, H_YEAR
]

PRIORITY_HEADERS = set()
for item in raw_priority_headers:
    if isinstance(item, (tuple, list, set)):
        PRIORITY_HEADERS.update(item)
    else:
        PRIORITY_HEADERS.add(item)

ENGLISH_PRIORITY_KEYS = {
    'Letter Type',
    'Student ID (e.g. 6587999)',
    'Prefix',
    'First Name',
    'Last Name',
    'Company/Organization Name for Internship',
    'Letter Addressed To (Company Contact Person)',
    'Position',
    'Internship Position',
    'Internship Period',
    'Year',
    'ชั้นปีที่',
}

ALL_PRIORITY_KEYS = PRIORITY_HEADERS.union(ENGLISH_PRIORITY_KEYS)
EXCLUDE_KEYS = {'id', 'batch', 'row_index', 'supported', 'missing', 'batch_id'}

THAI_MONTHS = {
    'มกราคม': 1, 'ม.ค': 1, 'กุมภาพันธ์': 2, 'ก.พ': 2, 'มีนาคม': 3, 'มี.ค': 3,
    'เมษายน': 4, 'เม.ย': 4, 'พฤษภาคม': 5, 'พ.ค': 5, 'มิถุนายน': 6, 'มิ.ย': 6,
    'กรกฎาคม': 7, 'ก.ค': 7, 'สิงหาคม': 8, 'ส.ค': 8, 'กันยายน': 9, 'ก.ย': 9,
    'ตุลาคม': 10, 'ต.ค': 10, 'พฤศจิกายน': 11, 'พ.ย': 11, 'ธันวาคม': 12, 'ธ.ค': 12,
}


# ==========================================
# HELPER FUNCTIONS
# ==========================================

def _pick(data, keys):
    for k in keys:
        v = data.get(k)
        if v and str(v).strip():
            return str(v).strip()
    return ''


def _field_ready(field):
    """Whether a FileField/FieldFile actually has a stored file, checked in a way
    that works regardless of storage backend (local filesystem, S3, etc). Using
    field.path + os.path.exists only works for filesystem storage and raises
    NotImplementedError on backends (e.g. S3) that don't support absolute paths."""
    if not field:
        return False
    try:
        return field.storage.exists(field.name)
    except Exception:
        return False


def log_activity(request, action, details=""):
    if not request.user or not request.user.is_authenticated:
        return

    x_forwarded_for = request.META.get('HTTP_X_FORWARDED_FOR')
    ip = x_forwarded_for.split(',')[0].strip() if x_forwarded_for else request.META.get('REMOTE_ADDR')

    AuditLog.objects.create(
        user=request.user,
        action=action,
        details=details,
        ip_address=ip
    )


def _load_table(uploaded_file):
    name = uploaded_file.name.lower()
    if name.endswith('.csv'):
        return pd.read_csv(uploaded_file, dtype=str).fillna('')
    return pd.read_excel(uploaded_file, dtype=str).fillna('')


def _clean_str(s):
    if not s:
        return ''
    s = str(s).lower()
    prefixes = ['mr.', 'mr', 'mrs.', 'mrs', 'ms.', 'ms', 'miss', 'dr.', 'dr', 'นาย', 'นาง', 'นางสาว']
    for p in prefixes:
        if s.startswith(p):
            s = s[len(p):]
    return re.sub(r'[^a-zA-Z0-9\u0e00-\u0e7f]', '', s)


def get_student_email(record):
    data = record.data or {}

    def norm(s):
        return re.sub(r'[^a-zA-Z0-9\u0e00-\u0e7f]', '', str(s)).lower()

    student_keys = ['studentprimaryemail', 'studentemail', 'studentmail', 'อีเมลนักศึกษา', 'student']
    company_keys = ['company', 'contractor', 'contact', 'hr', 'supervisor', 'บริษัท', 'ผู้ติดต่อ', 'ผู้ประสานงาน']

    for k, v in data.items():
        if not v or '@' not in str(v):
            continue
        nk = norm(k)
        if any(sk in nk for sk in student_keys) and not any(ck in nk for ck in company_keys):
            return str(v).strip()

    for k, v in data.items():
        if not v or '@' not in str(v):
            continue
        val_str = str(v).strip().lower()
        nk = norm(k)
        if any(dom in val_str for dom in ['@student', '@u.mahidol', '@mail.mahidol', '.ac.th', '.edu']):
            if not any(ck in nk for ck in company_keys):
                return str(v).strip()

    for k, v in data.items():
        if not v or '@' not in str(v):
            continue
        nk = norm(k)
        if not any(ck in nk for ck in company_keys):
            return str(v).strip()

    return ''


def _get_letter_type_keywords(letter_type_raw):
    category = normalize_letter_type(letter_type_raw)
    keywords = set()

    if category == "request":
        keywords.update([
            _clean_str('Request Letter for Internship Placement'),
            _clean_str('หนังสือขอความอนุเคราะห์'),
            _clean_str('ขอความอนุเคราะห์'),
            'request', 'req'
        ])
    elif category == "confirm":
        keywords.update([
            _clean_str('Student Referral Letter'),
            _clean_str('หนังสือส่งตัวนักศึกษาเข้าฝึกงาน'),
            _clean_str('หนังสือส่งตัว'),
            'referral', 'confirm', 'send', 'dispatch'
        ])
    elif category == "extension":
        keywords.update([
            _clean_str('หนังสือขอขยายระยะเวลาฝึกงาน'),
            _clean_str('ขยายระยะเวลา'),
            'extension', 'ext'
        ])
    else:
        keywords.add(_clean_str(letter_type_raw))

    return {k for k in keywords if k}


def _match_and_save_pdf(records, filename, file_content):
    clean_filename = _clean_str(filename)
    candidates = []

    for record in records:
        clean_id = _clean_str(record.student_id)
        clean_name = _clean_str(record.display_name)

        matched = (clean_id and clean_id in clean_filename) or \
                  (clean_name and len(clean_name) >= 3 and clean_name in clean_filename)
        if matched:
            candidates.append(record)

    if not candidates:
        return False

    if len(candidates) == 1:
        record = candidates[0]
    else:
        type_matches = [
            r for r in candidates
            if _get_letter_type_keywords(r.letter_type) & {
                clean_filename[i:j]
                for i in range(len(clean_filename)) for j in range(i + 1, len(clean_filename) + 1)
            }
        ]
        if len(type_matches) == 1:
            record = type_matches[0]
        else:
            return False

    doc = GeneratedDocument.objects.filter(record=record).last()
    if not doc:
        ref_no = str(2600 + record.row_index)
        doc = GeneratedDocument.objects.create(
            record=record,
            ref_no=ref_no,
            status=GeneratedDocument.STATUS_GENERATED,
            file_name=filename,
        )
    else:
        doc.status = GeneratedDocument.STATUS_GENERATED

    doc.signed_file.save(os.path.basename(filename), file_content, save=True)
    return True


def _get_val(row, *keys):
    for key in keys:
        val = row.get(key)
        if val is None or pd.isna(val):
            continue
        if isinstance(val, float) and val.is_integer():
            val = int(val)

        val_str = str(val).strip()
        if val_str.endswith('.0') and val_str[:-2].isdigit():
            val_str = val_str[:-2]
        if val_str.lower() in ('nan', 'none', 'nat'):
            continue
        if val_str:
            return val_str
    return ''


def _parse_thai_date(date_str):
    date_str = date_str.strip()
    m = re.match(r'(\d{1,2})\s+([^\d]+?)\.?\s*(\d{2,4})?$', date_str)
    if not m:
        return None

    day = int(m.group(1))
    month_raw = m.group(2).strip().rstrip('.')
    month = THAI_MONTHS.get(month_raw)
    if not month:
        return None

    year_raw = m.group(3)
    year = (int(year_raw) - 543) if year_raw else None
    return day, month, year


def _get_semester_from_period(period_str, default='2026/1'):
    if not period_str:
        return default

    range_match = re.search(r'\((.*?)\)', period_str)
    if not range_match:
        return default

    parts = range_match.group(1).split('-')
    if len(parts) != 2:
        return default

    end = _parse_thai_date(parts[1])
    if not end or end[2] is None:
        return default
    end_year = end[2]

    start = _parse_thai_date(parts[0])
    if start and start[1] is not None:
        start_month = start[1]
    else:
        start_month = end[1]

    if start_month in (8, 9, 10, 11, 12):
        sem = 1
    elif start_month in (1, 2, 3, 4, 5):
        sem = 2
    else:
        sem = 3

    return f"{end_year}/{sem}"


def _check_year_eligibility(record):
    if not record.period:
        return False
    m = re.search(r'ปี\s*(\d+)\s*(?:-\s*(\d+))?\s*เท่านั้น', record.period)
    if not m:
        return False

    start = int(m.group(1))
    end = int(m.group(2)) if m.group(2) else start
    allowed = set(range(start, end + 1))

    digits = re.sub(r'\D', '', record.year or '')
    if not digits:
        return None

    return int(digits) not in allowed


# ==========================================
# AUTHENTICATION & PROFILE VIEWS
# ==========================================

def login_view(request):
    if request.user.is_authenticated:
        return redirect('letters:index')

    if request.method == 'POST':
        form = AuthenticationForm(request, data=request.POST)
        if form.is_valid():
            user = form.get_user()
            if not user.is_staff:
                messages.error(request, "This account does not have staff permissions.")
                return render(request, 'letters/login.html', {'form': form})
            login(request, user)
            request.session.set_expiry(1800)
            next_url = request.GET.get('next')
            return redirect(next_url if next_url else 'letters:index')
        else:
            messages.error(request, "Invalid username or password.")
    else:
        form = AuthenticationForm()
    return render(request, 'letters/login.html', {'form': form})


def logout_view(request):
    logout(request)
    return redirect('letters:login')


@staff_required
def profile_view(request):
    password_form = PasswordChangeForm(request.user)

    if request.method == 'POST':
        action = request.POST.get('action')

        if action == 'update_info':
            request.user.email = request.POST.get('email', '').strip()
            request.user.first_name = request.POST.get('first_name', '').strip()
            request.user.last_name = request.POST.get('last_name', '').strip()
            request.user.save()
            messages.success(request, "Profile updated successfully.")
            return redirect('letters:profile')

        elif action == 'change_password':
            password_form = PasswordChangeForm(request.user, request.POST)
            if password_form.is_valid():
                user = password_form.save()
                update_session_auth_hash(request, user)
                messages.success(request, "Password updated successfully.")
                return redirect('letters:profile')
            else:
                messages.error(request, "Please fix the password errors below.")

    return render(request, 'letters/profile.html', {'password_form': password_form})


# ==========================================
# CORE SYSTEM VIEWS
# ==========================================

@staff_required
def index(request):
    recent = UploadedBatch.objects.order_by('-uploaded_at')[:5]
    return render(request, 'letters/index.html', {'recent': recent})


@staff_required
def upload(request):
    if request.method != 'POST' or 'datafile' not in request.FILES:
        messages.error(request, "Please select a valid file to upload.")
        return redirect('letters:index')

    f = request.FILES['datafile']
    try:
        df = _load_table(f)
    except Exception as e:
        messages.error(request, f"Failed to parse file: {e}")
        return redirect('letters:index')

    batch = UploadedBatch.objects.create(file_name=f.name, uploaded_by=request.user)
    records = []
    in_memory_counts = {}

    for i, row in enumerate(df.to_dict(orient='records')):
        s_id_raw = _get_val(row, *H_STUDENT_ID)
        if not s_id_raw:
            continue

        s_id = str(s_id_raw).replace('.0', '').strip()
        if not s_id:
            continue

        l_type = _get_val(row, *H_LETTER_TYPE)
        course = _get_val(row, *H_COURSE)
        year_val = _get_val(row, *H_YEAR)
        company_name = _get_val(row, *H_COMPANY)
        contractor_name = _get_val(row, *H_CONTACT_PERSON)
        contractor_pos = _get_val(row, *H_CONTACT_POSITION)
        intern_pos = _get_val(row, *H_INTERN_POSITION)
        period_val = _get_val(row, *H_PERIOD)
        sem_val = _get_semester_from_period(period_val)

        prefix = _get_val(row, *H_PREFIX)
        first_name = _get_val(row, *H_FIRST_NAME)
        last_name = _get_val(row, *H_LAST_NAME)
        full_name = f"{prefix} {first_name} {last_name}".strip()

        combo_key = (s_id, l_type, course, sem_val)

        if combo_key not in in_memory_counts:
            db_count = StudentRecord.objects.filter(
                student_id=s_id,
                letter_type=l_type,
                course_id=course,
                semester=sem_val
            ).count()
            in_memory_counts[combo_key] = db_count

        in_memory_counts[combo_key] += 1
        times_val = in_memory_counts[combo_key]

        records.append(StudentRecord(
            batch=batch,
            row_index=i,
            student_id=s_id,
            display_name=full_name,
            letter_type=l_type,
            course_id=course,
            semester=sem_val,
            times=times_val,
            company=company_name,
            contractor=contractor_name,
            contractor_position=contractor_pos,
            internship_position=intern_pos,
            year=year_val,
            period=period_val,
            data=row,
        ))

    if records:
        StudentRecord.objects.bulk_create(records)
        messages.success(request, f"Successfully uploaded and extracted {len(records)} record(s).")
        return redirect('letters:preview', batch_id=batch.id)

    batch.delete()
    messages.error(request, "No valid student records found. Please check column header names.")
    return redirect('letters:index')


@staff_required
def preview(request, batch_id):
    batch = get_object_or_404(UploadedBatch, id=batch_id)
    records = list(batch.records.all())
    for r in records:
        r.supported = r.letter_type.strip() in SUPPORTED_LETTER_TYPES
        r.missing = [f for f in IMPORTANT_FIELDS if not getattr(r, f).strip()]
        r.year_flag = _check_year_eligibility(r)

    counts = {}
    for r in records:
        counts[r.letter_type] = counts.get(r.letter_type, 0) + 1

    return render(request, 'letters/preview.html', {
        'batch': batch,
        'records': records,
        'counts': counts,
    })


@staff_required
def record_detail(request, record_id):
    record = get_object_or_404(StudentRecord, id=record_id)
    document = GeneratedDocument.objects.filter(record=record).last()

    if request.method == 'POST':
        updated_data = dict(record.data or {})
        for field_key in list(updated_data.keys()):
            if field_key in request.POST:
                updated_data[field_key] = request.POST.get(field_key)

        record.data = updated_data
        record.company = _pick(updated_data, H_COMPANY) or record.company
        record.contractor = _pick(updated_data, H_CONTACT_PERSON) or record.contractor
        record.contractor_position = _pick(updated_data, H_CONTACT_POSITION) or record.contractor_position
        record.internship_position = _pick(updated_data, H_INTERN_POSITION) or record.internship_position
        record.period = _pick(updated_data, H_PERIOD) or record.period
        record.year = _pick(updated_data, H_YEAR) or record.year
        record.save()

        ref_no = document.ref_no if (document and document.ref_no) else str(2600 + record.row_index)
        buf, fname_or_reason = generate_for_row(record.data, ref_no=ref_no)
        if buf:
            if not document:
                document = GeneratedDocument(record=record, ref_no=ref_no)
            document.status = GeneratedDocument.STATUS_GENERATED
            document.file_name = fname_or_reason
            document.file.save(fname_or_reason, ContentFile(buf.getvalue()), save=True)

        log_activity(request, 'EDIT_RECORD', f"Updated student record #{record.id} ({record.display_name})")
        return redirect('letters:record_detail', record_id=record.id)

    if not document or not document.file:
        ref_no = str(2600 + record.row_index)
        buf, fname_or_reason = generate_for_row(record.data, ref_no=ref_no)
        if buf:
            if not document:
                document = GeneratedDocument.objects.create(
                    record=record,
                    ref_no=ref_no,
                    status=GeneratedDocument.STATUS_GENERATED,
                    file_name=fname_or_reason,
                )
            document.file.save(fname_or_reason, ContentFile(buf.getvalue()), save=True)

    raw_data = record.data or {}
    valid_keys = [
        k for k in raw_data.keys()
        if k and not k.startswith('Unnamed') and k.lower() not in EXCLUDE_KEYS
    ]

    primary_fields = [(k, raw_data[k]) for k in valid_keys if k in ALL_PRIORITY_KEYS]
    secondary_fields = [(k, raw_data[k]) for k in valid_keys if k not in ALL_PRIORITY_KEYS]

    return render(request, 'letters/record_detail.html', {
        'record': record,
        'primary_fields': primary_fields,
        'secondary_fields': secondary_fields,
        'document': document,
    })


@staff_required
def add_record(request):
    if request.method == 'POST':
        student_id = request.POST.get('student_id', '').strip()
        display_name = request.POST.get('display_name', '').strip()
        letter_type = request.POST.get('letter_type', '').strip()
        year = request.POST.get('year', '').strip()
        course_id = request.POST.get('course_id', '').strip()
        semester = request.POST.get('semester', '').strip()
        company = request.POST.get('company', '').strip()
        contractor = request.POST.get('contractor', '').strip()
        internship_position = request.POST.get('internship_position', '').strip()
        period = request.POST.get('period', '').strip()

        existing_attempts = StudentRecord.objects.filter(
            student_id=student_id,
            letter_type=letter_type,
            course_id=course_id,
            semester=semester
        ).count()

        batch = UploadedBatch.objects.order_by('-uploaded_at').first()
        if not batch:
            batch = UploadedBatch.objects.create(file_name="Manual Entry", uploaded_by=request.user)

        new_record = StudentRecord.objects.create(
            batch=batch,
            student_id=student_id,
            display_name=display_name,
            letter_type=letter_type,
            course_id=course_id,
            semester=semester,
            year=year,
            times=existing_attempts + 1,
            company=company,
            contractor=contractor,
            internship_position=internship_position,
            period=period,
            data={
                'Student ID (e.g. 6587999)': student_id,
                'First Name': display_name,
                'ชั้นปีที่': year,
                'Letter Type': letter_type,
                'Course ID': course_id,
                'Semester': semester,
                'Company/Organization Name for Internship': company,
                'Letter Addressed To (Company Contact Person)': contractor,
                'Internship Position': internship_position,
                'Internship Period': period,
            }
        )

        log_activity(request, 'ADD_RECORD', f"Manually added record for {student_id} ({course_id} {semester})")
        messages.success(request, f"Added letter request #{new_record.times} for {display_name}.")
        return redirect('letters:record_detail', record_id=new_record.id)

    return render(request, 'letters/add_record.html')


@staff_required
def audit_logs_view(request):
    logs = AuditLog.objects.select_related('user').all()[:100]
    return render(request, 'letters/audit_logs.html', {'logs': logs})


logger = logging.getLogger(__name__)
@staff_required
def generate(request, batch_id):
    batch = get_object_or_404(UploadedBatch, id=batch_id)

    if request.method == "POST":
        selected_ids = request.POST.getlist("record_id")
        fmt = request.POST.get("format", "docx").lower()

        if not selected_ids:
            messages.error(request, "No records selected.")
            return redirect("letters:preview", batch_id=batch_id)

        records = StudentRecord.objects.filter(id__in=selected_ids, batch=batch)
        zip_buffer = io.BytesIO()
        added_count = 0
        errors = []

        with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zf:
            if fmt == "docx":
                # DOCX generation is purely in-memory and fast
                for record in records:
                    try:
                        buf, filename = generate_for_row(record.data, ref_no="___")
                        if buf:
                            zf.writestr(filename, buf.getvalue())
                            added_count += 1
                    except Exception as e:
                        errors.append(f"Record #{record.id}: {str(e)}")
            else:
                # PDF Format: Single-pass batch conversion to prevent worker timeout/OOM
                with tempfile.TemporaryDirectory() as tmpdir:
                    file_map = {}  # doc_filename -> (zip_output_filename, generated_pdf_path)
                    docx_paths = []

                    # 1. Generate and save all modified DOCX files to tmpdir
                    for record in records:
                        try:
                            ref = str(record.student_id or record.id or "___")
                            res = generate_for_row(record.data, ref_no=ref)
                            if not res or res[0] is None:
                                continue

                            docx_buf, orig_fname = res
                            pdf_fname = orig_fname.rsplit(".", 1)[0] + ".pdf"

                            docx_buf.seek(0)
                            pdf_input_doc = Document(docx_buf)
                            apply_thai_font_fix(pdf_input_doc, font_name=PDF_SAFE_FONT)

                            doc_filename = f"doc_{record.id}.docx"
                            input_docx_path = os.path.join(tmpdir, doc_filename)
                            pdf_input_doc.save(input_docx_path)

                            docx_paths.append(input_docx_path)
                            file_map[doc_filename] = (
                                pdf_fname,
                                os.path.join(tmpdir, f"doc_{record.id}.pdf"),
                            )
                        except Exception as e:
                            errors.append(f"Record #{record.id}: {str(e)}")

                    # 2. Batch-convert ALL documents in ONE LibreOffice process
                    if docx_paths:
                        soffice_bin = _get_libreoffice_bin()
                        if not soffice_bin:
                            messages.error(request, "ไม่พบโปรแกรม LibreOffice ในระบบ")
                            return redirect("letters:preview", batch_id=batch_id)

                        env = os.environ.copy()
                        env["LANG"] = "th_TH.UTF-8"
                        env["LC_ALL"] = "th_TH.UTF-8"

                        lo_profile_dir = os.path.join(tmpdir, "lo_profile")
                        os.makedirs(lo_profile_dir, exist_ok=True)
                        profile_uri = Path(lo_profile_dir).resolve().as_uri()

                        cmd = [
                            soffice_bin,
                            "--headless",
                            "--invisible",
                            "--nologo",
                            "--norestore",
                            "--nofirststartwizard",
                            f"-env:UserInstallation={profile_uri}",
                            "--convert-to",
                            "pdf:writer_pdf_Export",
                            "--outdir",
                            tmpdir,
                            *docx_paths,  # Pass all file paths at once
                        ]

                        try:
                            subprocess.run(
                                cmd,
                                stdout=subprocess.PIPE,
                                stderr=subprocess.PIPE,
                                env=env,
                                timeout=90,
                            )
                        except subprocess.TimeoutExpired:
                            logger.error("Batch LibreOffice conversion timed out.")

                    # 3. Read converted PDFs into the ZIP archive
                    for doc_filename, (out_zip_fname, target_pdf_path) in file_map.items():
                        if os.path.exists(target_pdf_path):
                            with open(target_pdf_path, "rb") as f:
                                zf.writestr(out_zip_fname, f.read())
                                added_count += 1
                        else:
                            errors.append(f"PDF generation failed for {out_zip_fname}")

        if added_count == 0:
            first_err = errors[0] if errors else "PDF conversion failed."
            messages.error(request, f"Failed to generate documents. Reason: {first_err}")
            return redirect("letters:preview", batch_id=batch_id)

        zip_buffer.seek(0)
        response = HttpResponse(zip_buffer.getvalue(), content_type="application/zip")
        response["Content-Disposition"] = (
            f'attachment; filename="letters_batch_{batch.id}_{fmt}.zip"'
        )
        return response

    return redirect("letters:preview", batch_id=batch_id)


@staff_required
def download(request, run_id):
    gen_batch = get_object_or_404(GenerationBatch, id=run_id)
    generated = gen_batch.documents.filter(status=GeneratedDocument.STATUS_GENERATED)
    if not generated.exists():
        raise Http404("No documents in this batch")

    zip_buf = io.BytesIO()
    with zipfile.ZipFile(zip_buf, 'w', zipfile.ZIP_DEFLATED) as zf:
        for doc in generated:
            doc.file.open('rb')
            zf.writestr(doc.file_name, doc.file.read())
            doc.file.close()
    zip_buf.seek(0)
    return FileResponse(zip_buf, as_attachment=True, filename='internship_letters.zip')


# ==========================================
# FILE PREVIEW & DOWNLOAD VIEWS
# ==========================================

@staff_required
def download_record_pdf(request, record_id):
    record = get_object_or_404(StudentRecord, id=record_id)
    document = GeneratedDocument.objects.filter(record=record).last()

    if document and _field_ready(document.signed_file):
        document.signed_file.open('rb')
        response = FileResponse(document.signed_file, content_type='application/pdf')
        filename = os.path.basename(document.signed_file.name)
    else:
        ref_no = request.GET.get('ref_no', '___')
        pdf_buf, filename = generate_pdf_for_row(record.data, ref_no=ref_no)
        response = HttpResponse(pdf_buf.getvalue(), content_type='application/pdf')

    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


@xframe_options_sameorigin  # 1. Allows template to embed in <iframe>
@staff_required
def preview_record_pdf(request, record_id):
    record = get_object_or_404(StudentRecord, id=record_id)
    doc = GeneratedDocument.objects.filter(record=record).last()

    if not doc or not getattr(doc, 'file', None):
        return HttpResponse("No generated document found for this record.", status=404)

    # Use doc.file.open() instead of doc.file.path to support remote/cloud storage
    try:
        with doc.file.open('rb') as f:
            file_bytes = f.read()
    except Exception as e:
        return HttpResponse(f"Unable to read file from storage: {str(e)}", status=404)

    filename = os.path.basename(doc.file.name)

    # 1. Stream directly if file is already a PDF
    if doc.file.name.lower().endswith('.pdf'):
        response = HttpResponse(file_bytes, content_type='application/pdf')
        response['Content-Disposition'] = f'inline; filename="{filename}"'
        return response

    # 2. Convert DOCX to PDF stream via LibreOffice helper
    try:
        pdf_bytes = convert_docx_bytes_to_pdf(file_bytes)
        response = HttpResponse(pdf_bytes, content_type='application/pdf')
        response['Content-Disposition'] = f'inline; filename="student_{record.student_id}.pdf"'
        return response
    except Exception as e:
        return HttpResponse(f"PDF Conversion Failed: {str(e)}", status=500)




@staff_required
def download_letter_view(request, record_id, letter_type):
    record = get_object_or_404(StudentRecord, id=record_id)
    if letter_type not in ('confirm', 'request', 'extension'):
        raise Http404(f"Unknown letter type: '{letter_type}'")

    ref_no = str(2600 + record.row_index)
    buf, fname_or_reason = generate_for_row(record.data, ref_no=ref_no, kind_override=letter_type)
    if buf is None:
        raise Http404(fname_or_reason)

    buf.seek(0)
    return FileResponse(
        buf, as_attachment=True, filename=fname_or_reason,
        content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
    )


# ==========================================
# BATCH DISPATCH & SIGNED FILE MANAGEMENT
# ==========================================

@staff_required
def upload_signed(request, batch_id=None):
    selected_batch_id = request.GET.get('batch_id') or (str(batch_id) if batch_id else 'all')

    if request.method == 'POST':
        uploaded_files = request.FILES.getlist('signed_files') or request.FILES.getlist('files')
        matched_count = 0

        if selected_batch_id != 'all' and selected_batch_id.isdigit():
            batch = get_object_or_404(UploadedBatch, id=int(selected_batch_id))
            records = batch.records.all()
        else:
            all_records = StudentRecord.objects.order_by('-id')
            seen = set()
            records = []
            for r in all_records:
                key = (r.student_id.strip(), r.letter_type.strip(), r.times)
                if key not in seen:
                    seen.add(key)
                    records.append(r)

        for uploaded_file in uploaded_files:
            fname = uploaded_file.name.lower()
            if fname.endswith('.zip'):
                with zipfile.ZipFile(uploaded_file, 'r') as z:
                    for inner_name in z.namelist():
                        if inner_name.lower().endswith('.pdf') and not inner_name.startswith('__MACOSX'):
                            pdf_bytes = z.read(inner_name)
                            if _match_and_save_pdf(records, inner_name, ContentFile(pdf_bytes, name=inner_name)):
                                matched_count += 1
            elif fname.endswith('.pdf'):
                if _match_and_save_pdf(records, uploaded_file.name, uploaded_file):
                    matched_count += 1

        log_activity(
            request,
            'UPLOAD_SIGNED',
            f"Uploaded signed documents ({selected_batch_id}). Matched {matched_count} file(s)."
        )

        if matched_count > 0:
            messages.success(request, f"Successfully matched {matched_count} signed document(s).")
        else:
            messages.warning(request, "No matching students found. Ensure PDF filenames include Student IDs or Names.")

    return redirect(f"{reverse('letters:signed_preview')}?batch_id={selected_batch_id}")


@staff_required
def signed_preview(request, batch_id=None):
    batches = UploadedBatch.objects.order_by('-uploaded_at')
    selected_batch_id = request.GET.get('batch_id') or (str(batch_id) if batch_id else 'all')

    if selected_batch_id != 'all' and selected_batch_id.isdigit():
        records = StudentRecord.objects.filter(batch_id=int(selected_batch_id))
    else:
        records = StudentRecord.objects.select_related('batch').order_by('-id')

    # Fix N+1 Query: Fetch all relevant documents in ONE query and map to record_id
    docs = GeneratedDocument.objects.filter(record__in=records).order_by('record_id', 'id')
    latest_doc_map = {doc.record_id: doc for doc in docs}

    record_summary = []
    for r in records:
        doc = latest_doc_map.get(r.id)

        # Safe student email extraction
        try:
            student_email = get_student_email(r)
        except Exception:
            student_email = ""

        # Safe file readiness check (prevents crashes if file is missing from disk)
        has_signed = False
        if doc and getattr(doc, 'signed_file', None):
            try:
                has_signed = bool(_field_ready(doc.signed_file))
            except Exception:
                has_signed = False

        record_summary.append({
            'record': r,
            'doc': doc,
            'email': student_email,
            'has_signed': has_signed,
        })

    return render(request, 'letters/signed_preview.html', {
        'batches': batches,
        'selected_batch_id': str(selected_batch_id),
        'summary': record_summary,
    })

@staff_required
def send_signed_emails(request):
    if request.method != 'POST':
        return redirect('letters:signed_preview')

    selected_ids = request.POST.getlist('selected_records')
    selected_batch = request.POST.get('batch_id', 'all')

    records = StudentRecord.objects.filter(id__in=selected_ids)

    if not records.exists():
        messages.warning(request, "No student records were selected. Please check at least one box.")
        return redirect(f"{reverse('letters:signed_preview')}?batch_id={selected_batch}")

    staff_email = request.user.email if request.user.email else settings.DEFAULT_FROM_EMAIL
    staff_name = request.user.get_full_name() or request.user.username

    sent_count = 0
    failed = []

    for r in records:
        student_email = request.POST.get(f'email_{r.id}', '').strip() or get_student_email(r)
        doc = GeneratedDocument.objects.filter(record=r).last()

        if doc and _field_ready(doc.signed_file) and student_email:
            subject = f"Signed Internship Letter & Draft - {r.display_name}"
            body = (
                f"Dear {r.display_name},\n\n"
                f"Please find attached your official signed internship letter (PDF) "
                f"and draft copy.\n\n"
                f"Best regards,\n"
                f"{staff_name}\n"
                f"Faculty of ICT, Mahidol University\n"
                f"Email: {staff_email}"
            )

            email = EmailMessage(
                subject=subject,
                body=body,
                from_email=staff_email,
                to=[student_email],
                reply_to=[staff_email],
            )

            with doc.signed_file.open('rb') as f:
                email.attach(os.path.basename(doc.signed_file.name), f.read(), 'application/pdf')

            if _field_ready(doc.file):
                with doc.file.open('rb') as f:
                    is_pdf = doc.file.name.lower().endswith('.pdf')
                    c_type = 'application/pdf' if is_pdf else 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
                    email.attach(os.path.basename(doc.file.name), f.read(), c_type)

            try:
                email.send(fail_silently=False)
            except Exception as e:
                failed.append(f"{r.display_name} ({student_email}): {e}")
                continue

            doc.status = GeneratedDocument.STATUS_SENT
            doc.email_sent_at = timezone.now()
            doc.save()
            sent_count += 1

    log_activity(
        request,
        'SEND_EMAIL',
        f"Sent {sent_count} signed internship document email(s)." + (f" {len(failed)} failed." if failed else "")
    )

    if sent_count > 0:
        messages.success(request, f"Successfully sent {sent_count} email(s) to student(s).")
    if failed:
        messages.warning(request, f"{len(failed)} email(s) failed: " + "; ".join(failed[:5]) + (" …" if len(failed) > 5 else ""))
    if sent_count == 0 and not failed:
        messages.warning(request, "No emails sent. Verify students have attached signed documents and valid email addresses.")

    return redirect(f"{reverse('letters:signed_preview')}?batch_id={selected_batch}")


@staff_required
def view_signed_pdf(request, document_id):
    document = get_object_or_404(GeneratedDocument, id=document_id)
    if not _field_ready(document.signed_file):
        raise Http404("No signed file uploaded for this document yet.")

    document.signed_file.open('rb')
    response = FileResponse(document.signed_file, content_type='application/pdf')
    response['Content-Disposition'] = f'inline; filename="{os.path.basename(document.signed_file.name)}"'
    return response


@staff_required
def upload_signed_by_id(request, letter_id):
    if request.method == 'POST' and 'signed_file' in request.FILES:
        letter = get_object_or_404(StudentLetter, pk=letter_id)
        letter.signed_file = request.FILES['signed_file']
        letter.status = 'signed'
        letter.save()
        messages.success(request, f"Signed file uploaded for student {letter.student_id}.")
    else:
        messages.error(request, "Invalid file upload attempt.")

    return redirect('letters:index')


# ==========================================
# TEMPLATE & LETTER TYPE MANAGEMENT
# ==========================================

@staff_required
def template_list(request):
    templates = LetterTemplate.objects.order_by('letter_type', '-updated_at')
    for t in templates:
        mappings = list(t.field_mappings.all())
        t.total_fields = len(mappings)
        t.mapped_fields = len([m for m in mappings if m.source_type != TemplateFieldMapping.SOURCE_UNMAPPED])

    letter_types = LetterType.objects.order_by('display_name')
    return render(request, 'letters/template_list.html', {
        'templates': templates,
        'letter_types': letter_types,
    })


@staff_required
def template_upload(request):
    if request.method != 'POST' or 'template_file' not in request.FILES:
        messages.error(request, "Please select a template file and letter type.")
        return redirect('letters:template_list')

    letter_type = request.POST.get('letter_type', '').strip()
    if not LetterType.objects.filter(slug=letter_type).exists():
        messages.error(request, "Please select a valid, existing letter type.")
        return redirect('letters:template_list')

    f = request.FILES['template_file']
    template_obj, created = LetterTemplate.objects.update_or_create(
        letter_type=letter_type,
        defaults={'name': f.name, 'file': f, 'is_active': False},
    )
    template_obj.field_mappings.all().delete()

    template_obj.file.open('rb')
    try:
        detected = extract_fields_from_docx(template_obj.file)
    finally:
        template_obj.file.close()

    placeholders = detected.get('template_variables') or []
    if not placeholders:
        messages.warning(
            request,
            f"Uploaded '{f.name}', but no {{{{ tag }}}}-style placeholders were found in it."
        )
        return redirect('letters:template_list')

    for tag in placeholders:
        TemplateFieldMapping.objects.get_or_create(template=template_obj, placeholder=tag)

    action = 'Replaced' if not created else 'Uploaded'
    log_activity(request, 'UPLOAD_TEMPLATE', f"{action} template '{f.name}' for '{letter_type}'")
    messages.success(request, f"{action} '{f.name}'. Map its {len(placeholders)} field(s) before activating.")
    return redirect('letters:template_mapping', template_id=template_obj.id)


@staff_required
def lettertype_list(request):
    types = LetterType.objects.prefetch_related('mappings').order_by('display_name')
    unmapped_raw_values = (
        StudentRecord.objects.exclude(letter_type='')
        .values_list('letter_type', flat=True).distinct()
    )
    known_raw = set(LetterTypeMapping.objects.values_list('raw_text', flat=True))
    unmapped_raw_values = sorted({v.strip() for v in unmapped_raw_values if v.strip() not in known_raw})
    return render(request, 'letters/lettertype_list.html', {
        'types': types,
        'unmapped_raw_values': unmapped_raw_values,
    })


@staff_required
def lettertype_create(request):
    if request.method == 'POST':
        slug = request.POST.get('slug', '').strip()
        display_name = request.POST.get('display_name', '').strip()
        if not slug or not display_name:
            messages.error(request, "Both slug and display name are required.")
        else:
            LetterType.objects.get_or_create(slug=slug, defaults={'display_name': display_name})
            log_activity(request, 'CREATE_LETTER_TYPE', f"Created letter type '{slug}'")
            messages.success(request, f"Created letter type '{display_name}'.")
    return redirect('letters:lettertype_list')

@staff_required
def preview_record_docx(request, record_id):
    record = get_object_or_404(StudentRecord, id=record_id)
    doc = GeneratedDocument.objects.filter(record=record).last()

    if not doc or not getattr(doc, 'file', None):
        return HttpResponse("No generated document found for this record.", status=404)

    try:
        with doc.file.open('rb') as f:
            file_bytes = f.read()
        
        response = HttpResponse(
            file_bytes, 
            content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )
        filename = os.path.basename(doc.file.name)
        response['Content-Disposition'] = f'inline; filename="{filename}"'
        return response
    except Exception as e:
        return HttpResponse(f"Unable to read file: {str(e)}", status=404)

@staff_required
def mapping_create(request):
    if request.method == 'POST':
        raw_text = request.POST.get('raw_text', '').strip()
        letter_type_id = request.POST.get('letter_type_id')
        if not raw_text or not letter_type_id:
            messages.error(request, "Both the raw text and a letter type are required.")
        else:
            letter_type = get_object_or_404(LetterType, id=letter_type_id)
            LetterTypeMapping.objects.update_or_create(raw_text=raw_text, defaults={'letter_type': letter_type})
            log_activity(request, 'CREATE_TYPE_MAPPING', f"Mapped '{raw_text}' -> '{letter_type.slug}'")
            messages.success(request, f"Mapped '{raw_text}' to '{letter_type.display_name}'.")
    return redirect('letters:lettertype_list')


@staff_required
def template_mapping(request, template_id):
    template_obj = get_object_or_404(LetterTemplate, id=template_id)
    mappings = template_obj.field_mappings.order_by('placeholder')

    if request.method == 'POST':
        for mapping in mappings:
            mapping.source_type = request.POST.get(f'source_type_{mapping.id}', TemplateFieldMapping.SOURCE_UNMAPPED)
            mapping.source_value = request.POST.get(f'source_value_{mapping.id}', '').strip()
            mapping.save()
        log_activity(request, 'MAP_TEMPLATE_FIELDS', f"Updated mapping for template #{template_obj.id}")
        messages.success(request, "Field mapping saved.")
        return redirect('letters:template_mapping', template_id=template_obj.id)

    return render(request, 'letters/template_mapping.html', {
        'template': template_obj,
        'mappings': mappings,
        'available_fields': AVAILABLE_SOURCE_FIELDS,
    })


@staff_required
def template_preview(request, template_id):
    template_obj = get_object_or_404(LetterTemplate, id=template_id)
    sample_record = StudentRecord.objects.order_by('-id').first()
    if not sample_record:
        messages.error(request, "No student records exist yet to preview against.")
        return redirect('letters:template_mapping', template_id=template_obj.id)

    kind = normalize_letter_type(sample_record.letter_type)
    ref_no = str(2600 + sample_record.row_index)

    docx_buf = render_with_template(template_obj, sample_record, kind, ref_no)
    pdf_bytes = convert_docx_bytes_to_pdf(docx_buf.getvalue())

    response = HttpResponse(pdf_bytes, content_type='application/pdf')
    response['Content-Disposition'] = f'inline; filename="preview_{template_obj.name}.pdf"'
    return response


@staff_required
def template_activate(request, template_id):
    template_obj = get_object_or_404(LetterTemplate, id=template_id)
    unmapped = template_obj.field_mappings.filter(source_type=TemplateFieldMapping.SOURCE_UNMAPPED).count()
    if unmapped:
        messages.error(request, f"Can't activate — {unmapped} placeholder(s) still unmapped.")
        return redirect('letters:template_mapping', template_id=template_obj.id)

    template_obj.is_active = True
    template_obj.save()
    log_activity(request, 'ACTIVATE_TEMPLATE', f"Activated template #{template_obj.id} for '{template_obj.letter_type}'")
    messages.success(request, f"'{template_obj.name}' is now active for '{template_obj.letter_type}'.")
    return redirect('letters:template_list')